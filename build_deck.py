# -*- coding: utf-8 -*-
"""Generate TenderPilot AI investor pitch deck (20 slides, Arabic RTL)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- Brand palette ----------
BG_DARK   = RGBColor(0x0B, 0x1F, 0x2A)   # deep petrol blue
BG_PANEL  = RGBColor(0x12, 0x2E, 0x3D)
ACCENT    = RGBColor(0xD9, 0xB3, 0x6A)   # gold
ACCENT2   = RGBColor(0x3F, 0xC1, 0xB0)   # teal
WHITE     = RGBColor(0xF4, 0xF7, 0xF9)
MUTED     = RGBColor(0x9D, 0xB2, 0xBF)
GREEN     = RGBColor(0x4C, 0xC2, 0x7A)
RED       = RGBColor(0xE0, 0x6B, 0x6B)

FONT = "Tajawal"  # falls back gracefully; Arabic-friendly
FONT_FALLBACK = "Arial"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def set_rtl(para):
    pPr = para._pPr
    if pPr is None:
        pPr = para._p.get_or_add_pPr()
    pPr.set('rtl', '1')


def add_bg(slide, color=BG_DARK):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    slide.shapes._spTree.remove(s._element)
    slide.shapes._spTree.insert(2, s._element)
    return s


def rect(slide, x, y, w, h, color, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def txt(slide, x, y, w, h, runs, size=18, color=WHITE, bold=False,
        align=PP_ALIGN.RIGHT, rtl=True, anchor=MSO_ANCHOR.TOP, line_spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(2); tf.margin_bottom = Pt(2)
    if isinstance(runs, str):
        runs = [(runs, {})]
    first = True
    for line in runs:
        if isinstance(line, tuple):
            line = [line]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        if rtl:
            set_rtl(p)
        for seg, opt in line:
            r = p.add_run(); r.text = seg
            r.font.name = opt.get('font', FONT)
            r.font.size = Pt(opt.get('size', size))
            r.font.bold = opt.get('bold', bold)
            r.font.color.rgb = opt.get('color', color)
    return tb


def slide_number(slide, n):
    txt(slide, Inches(0.3), Inches(7.0), Inches(1.2), Inches(0.4),
        f"{n:02d} / 20", size=11, color=MUTED, align=PP_ALIGN.LEFT, rtl=False)
    txt(slide, Inches(11.3), Inches(7.0), Inches(1.8), Inches(0.4),
        "TenderPilot AI", size=11, color=ACCENT, align=PP_ALIGN.RIGHT, rtl=False)


def header(slide, kicker, title):
    # accent bar
    rect(slide, Inches(12.5), Inches(0.55), Inches(0.55), Inches(0.12), ACCENT)
    txt(slide, Inches(6.0), Inches(0.45), Inches(7.0), Inches(0.4),
        kicker, size=14, color=ACCENT, bold=True)
    txt(slide, Inches(5.0), Inches(0.85), Inches(8.0), Inches(0.9),
        title, size=30, color=WHITE, bold=True)


def new_slide(n, decorate=True):
    s = prs.slides.add_slide(BLANK)
    add_bg(s)
    if decorate:
        slide_number(s, n)
    return s


def bullets(slide, x, y, w, items, size=16, gap=0.62, icon="●", icon_color=ACCENT):
    cy = y
    for it in items:
        if isinstance(it, tuple):
            ic, text, icol = it
        else:
            ic, text, icol = icon, it, icon_color
        rect_h = Inches(0.5)
        txt(slide, x, Inches(cy), w, rect_h,
            [[(text + "  ", {'size': size, 'color': WHITE}), (ic, {'size': size, 'color': icol, 'bold': True})]],
            align=PP_ALIGN.RIGHT)
        cy += gap
    return cy


# ============ SLIDE 1 — Cover ============
s = new_slide(1, decorate=False)
# decorative gold corner
rect(s, Inches(0), Inches(0), Inches(0.18), SH, ACCENT)
txt(s, Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.2),
    "TenderPilot AI", size=58, color=WHITE, bold=True, align=PP_ALIGN.CENTER, rtl=False)
txt(s, Inches(1.0), Inches(3.5), Inches(11.3), Inches(0.8),
    "موظف المناقصات الحكومية الذكي — لكل شركة في الخليج", size=26, color=ACCENT, align=PP_ALIGN.CENTER)
txt(s, Inches(1.0), Inches(4.5), Inches(11.3), Inches(0.6),
    "AI-Native Government Tendering Platform  ·  Built for Saudi Vision 2030",
    size=16, color=MUTED, align=PP_ALIGN.CENTER, rtl=False)
txt(s, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.5),
    "Investor Pitch Deck  ·  Series Seed → Series A", size=13, color=MUTED, align=PP_ALIGN.CENTER, rtl=False)

# ============ SLIDE 2 — Problem ============
s = new_slide(2)
header(s, "المشكلة · THE PROBLEM", "سوق ضخم + عملية بدائية = فرص ضائعة بالمليارات")
items = [
    (("✕", "اكتشاف يدوي: الشركات لا تتابع آلاف المنافسات في اعتماد والجهات.", RED)),
    (("✕", "إعداد بطيء: العرض الفني يستغرق 2–6 أسابيع من فريق مكلف.", RED)),
    (("✕", "استبعاد شكلي: نسبة كبيرة من العروض تُرفض لأسباب امتثالية يمكن منعها.", RED)),
    (("✕", "قرار أعمى: لا أداة تحسب احتمالية الفوز قبل صرف الجهد.", RED)),
]
rect(s, Inches(0.7), Inches(2.1), Inches(11.9), Inches(3.4), BG_PANEL)
bullets(s, Inches(1.1), 2.55, Inches(11.1), items, size=19, gap=0.78)

# ============ SLIDE 3 — Why Now ============
s = new_slide(3)
header(s, "لماذا الآن · WHY NOW", "التقاء نادر بين الطلب والتقنية والتنظيم")
cards = [
    ("رؤية 2030", "إنفاق حكومي/شبه حكومي تاريخي ورقمنة المشتريات عبر اعتماد."),
    ("نضج وكلاء AI", "قدرة على الاستدلال على وثائق عربية طويلة ومعقّدة."),
    ("تنظيم محفّز", "PDPL ودفع نحو المحتوى المحلي والرقمنة."),
    ("فجوة سوقية", "لا لاعب SaaS متخصص بمستوى Enterprise إقليمياً."),
]
x = Inches(0.7); w = Inches(2.85); gap = Inches(0.18)
for i, (t, d) in enumerate(cards):
    cx = Emu(int(x) + i * (int(w) + int(gap)))
    rect(s, cx, Inches(2.4), w, Inches(3.0), BG_PANEL)
    rect(s, cx, Inches(2.4), w, Inches(0.12), ACCENT2)
    txt(s, Emu(int(cx)+Pt(6)), Inches(2.7), w, Inches(0.7), t, size=18, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Emu(int(cx)+Pt(6)), Inches(3.5), Emu(int(w)-int(Pt(12))), Inches(1.7), d, size=14, color=WHITE, align=PP_ALIGN.CENTER)

# ============ SLIDE 4 — Solution ============
s = new_slide(4)
header(s, "الحل · THE SOLUTION", "منصة AI واحدة تغطي رحلة المناقصة من الاكتشاف حتى التسليم")
flow = ["اكتشاف", "تأهيل", "تحليل", "كتابة", "قرار", "عقد", "تسليم"]
x0 = Inches(0.7); tot = Inches(11.9); n = len(flow)
cw = Emu(int(tot)//n - int(Pt(8)))
for i, st in enumerate(flow):
    cx = Emu(int(x0) + i*(int(tot)//n))
    rect(s, cx, Inches(2.5), cw, Inches(1.0), BG_PANEL)
    txt(s, cx, Inches(2.7), cw, Inches(0.6), st, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
modules = "Discovery · Opportunity Matching · Tender Analyzer · Bid Writer · Competitor Intelligence · Win Probability · Contract Intelligence · Delivery Readiness"
rect(s, Inches(0.7), Inches(4.1), Inches(11.9), Inches(1.6), BG_PANEL)
txt(s, Inches(1.0), Inches(4.35), Inches(11.3), Inches(0.5), "ثماني وحدات أساسية مدعومة بالذكاء الاصطناعي:", size=16, color=ACCENT, bold=True)
txt(s, Inches(1.0), Inches(4.85), Inches(11.3), Inches(0.8), modules, size=15, color=WHITE, align=PP_ALIGN.CENTER, rtl=False)
txt(s, Inches(0.7), Inches(6.1), Inches(11.9), Inches(0.5),
    "ليست ميزة AI مضافة — بل نظام تشغيل كامل لقسم تطوير الأعمال.", size=15, color=MUTED, align=PP_ALIGN.CENTER)

# ============ SLIDE 5 — Demo ============
s = new_slide(5)
header(s, "العرض الحي · PRODUCT", "من كراسة شروط إلى عرض احترافي في ساعات لا أسابيع")
demo = [
    ("1 · رفع الكراسة", "استخراج المتطلبات والمخاطر تلقائياً + Opportunity Score من 0–100."),
    ("2 · احتمالية الفوز", "Win Probability % مع تفسير شفّاف لكل العوامل المؤثرة."),
    ("3 · بناء العرض", "Proposal Builder يولّد العرض الفني مع فحص امتثال فوري."),
]
x = Inches(0.7); w = Inches(3.83); gap = Inches(0.2)
for i, (t, d) in enumerate(demo):
    cx = Emu(int(x) + i*(int(w)+int(gap)))
    rect(s, cx, Inches(2.4), w, Inches(3.4), BG_PANEL)
    rect(s, cx, Inches(2.4), w, Inches(0.6), BG_DARK)
    txt(s, cx, Inches(2.5), w, Inches(0.5), t, size=17, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Emu(int(cx)+int(Pt(8))), Inches(3.3), Emu(int(w)-int(Pt(16))), Inches(2.2), d, size=15, color=WHITE, align=PP_ALIGN.CENTER)

# ============ SLIDE 6 — Market ============
s = new_slide(6)
header(s, "حجم السوق · MARKET SIZE", "سوق كبير وقابل للتوسّع الخليجي")
rings = [
    ("TAM", "+250,000 منشأة موردة للحكومات في الخليج", Inches(4.2), BG_PANEL),
    ("SAM", "40,000–60,000 منشأة سعودية نشطة في المناقصات", Inches(3.0), BG_DARK),
    ("SOM", "2,500–4,000 عميل مدفوع خلال 3 سنوات", Inches(1.8), ACCENT2),
]
cx = Inches(3.2); cy = Inches(4.1)
for label, desc, dia, col in rings:
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx)-int(dia)//2), Emu(int(cy)-int(dia)//2), dia, dia)
    c.fill.solid(); c.fill.fore_color.rgb = col; c.line.color.rgb = ACCENT; c.line.width = Pt(1.5)
    c.shadow.inherit = False
for i, (label, desc, dia, col) in enumerate(rings):
    txt(s, Inches(6.6), Inches(2.4 + i*1.15), Inches(6.0), Inches(1.0),
        [[(desc, {'size': 15, 'color': WHITE}), ("   " + label + " ", {'size': 20, 'color': ACCENT, 'bold': True})]],
        align=PP_ALIGN.RIGHT)

# ============ SLIDE 7 — Agents ============
s = new_slide(7)
header(s, "المنتج · MULTI-AGENT ENGINE", "تسعة وكلاء متخصصين يعملون كفريق مناقصات كامل")
agents = ["Scout", "Analysis", "Proposal", "Risk", "Pricing", "Compliance", "Competitor", "Contract", "Executive"]
# orchestrator center
oc = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.66), Inches(3.5), Inches(2.0), Inches(1.4))
oc.fill.solid(); oc.fill.fore_color.rgb = ACCENT; oc.line.fill.background(); oc.shadow.inherit=False
txt(s, Inches(5.66), Inches(3.85), Inches(2.0), Inches(0.7), "Orchestrator", size=15, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER, rtl=False)
import math
cxc, cyc = 6.66, 4.2
for i, a in enumerate(agents):
    ang = math.radians(i*(360/9) - 90)
    px = cxc + 4.4*math.cos(ang) - 0.85
    py = cyc + 2.0*math.sin(ang) - 0.32
    rect(s, Inches(px), Inches(py), Inches(1.7), Inches(0.64), BG_PANEL, line=ACCENT2)
    txt(s, Inches(px), Inches(py+0.08), Inches(1.7), Inches(0.5), a, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, rtl=False)

# ============ SLIDE 8 — Moat ============
s = new_slide(8)
header(s, "الخندق الدفاعي · THE MOAT", "ميزة تتعمّق مع الوقت — صعبة التقليد")
moats = [
    ("Data Moat", "قاعدة معرفة للترسيات والمنافسين تكبر كل يوم."),
    ("Compliance Moat", "محرك امتثال مُحدّث وفق الأنظمة الحكومية."),
    ("Workflow Lock-in", "المنصة تصبح نظام التشغيل اليومي لتطوير الأعمال."),
    ("Integration Moat", "تكامل اعتماد + ERP/CRM للعميل."),
]
x = Inches(0.7); w = Inches(5.85)
for i, (t, d) in enumerate(moats):
    col = i % 2; row = i // 2
    cx = Emu(int(x) + col*(int(w)+int(Inches(0.2))))
    cy = Inches(2.4 + row*1.55)
    rect(s, cx, cy, w, Inches(1.35), BG_PANEL)
    rect(s, Emu(int(cx)+int(w)-int(Inches(0.12))), cy, Inches(0.12), Inches(1.35), ACCENT)
    txt(s, Emu(int(cx)+int(Pt(8))), Emu(int(cy)+int(Pt(8))), Emu(int(w)-int(Pt(20))), Inches(0.5), t, size=18, color=ACCENT, bold=True)
    txt(s, Emu(int(cx)+int(Pt(8))), Emu(int(cy)+int(Inches(0.6))), Emu(int(w)-int(Pt(20))), Inches(0.6), d, size=14, color=WHITE)

# ============ SLIDE 9 — Traction ============
s = new_slide(9)
header(s, "الجر والتحقق · TRACTION", "قيمة مُثبتة لا وعود")
kpis = [
    ("Pilots & Logos", "تجارب وعملاء أوائل"),
    ("Win-rate ↑", "زيادة % في معدل الفوز"),
    ("Time-to-Proposal", "من أسابيع → ساعات"),
    ("Pipeline / LOIs", "خط أنابيب مبيعات نشط"),
]
x = Inches(0.7); w = Inches(2.85); gap = Inches(0.18)
for i, (t, d) in enumerate(kpis):
    cx = Emu(int(x) + i*(int(w)+int(gap)))
    rect(s, cx, Inches(2.5), w, Inches(2.6), BG_PANEL)
    txt(s, cx, Inches(2.9), w, Inches(0.7), t, size=17, color=ACCENT, bold=True, align=PP_ALIGN.CENTER, rtl=False)
    txt(s, cx, Inches(3.8), w, Inches(1.0), d, size=14, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, Inches(0.7), Inches(5.5), Inches(11.9), Inches(0.6),
    "تُحدّث هذه الشريحة بأرقامك الفعلية (Pilots / LOIs / نتائج win-rate).", size=14, color=MUTED, align=PP_ALIGN.CENTER)

# ============ SLIDE 10 — Business Model ============
s = new_slide(10)
header(s, "نموذج العمل · BUSINESS MODEL", "اشتراك + استهلاك AI = إيراد متكرر عالي الهامش")
plans = [
    ("Starter", "شركات صغيرة", "9K–15K ريال"),
    ("Professional", "متوسطة / مقاولون", "30K–60K ريال"),
    ("Enterprise", "كبيرة / استشارات", "120K–400K+ ريال"),
    ("Government", "جهات حكومية", "تعاقدي مخصّص"),
]
x = Inches(0.7); w = Inches(2.85); gap = Inches(0.18)
for i, (t, seg, price) in enumerate(plans):
    cx = Emu(int(x) + i*(int(w)+int(gap)))
    rect(s, cx, Inches(2.4), w, Inches(3.2), BG_PANEL)
    rect(s, cx, Inches(2.4), w, Inches(0.7), ACCENT if i < 3 else ACCENT2)
    txt(s, cx, Inches(2.5), w, Inches(0.5), t, size=18, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER, rtl=False)
    txt(s, cx, Inches(3.4), w, Inches(0.6), seg, size=14, color=MUTED, align=PP_ALIGN.CENTER)
    txt(s, cx, Inches(4.5), w, Inches(0.7), price, size=18, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.7), Inches(5.9), Inches(11.9), Inches(0.5),
    "Land عبر Starter/Pro ثم Expand داخل المؤسسة نحو Enterprise/Government.", size=14, color=MUTED, align=PP_ALIGN.CENTER)

# ============ SLIDE 11 — GTM ============
s = new_slide(11)
header(s, "الذهاب للسوق · GTM", "ثلاث محركات نمو متكاملة")
gtm = [
    ("Product-Led", "تجربة ذاتية + SEO عربي + محتوى تعليمي يجذب الشركات الصغيرة والمتوسطة."),
    ("Sales-Led", "Inside + Field sales للـ Enterprise/Gov عبر POCs وعلاقات الجهات."),
    ("Partner-Led", "غرف تجارية، حاضنات، شركاء ERP، ومكاتب استشارية كقنوات توزيع."),
]
for i, (t, d) in enumerate(gtm):
    cy = Inches(2.4 + i*1.15)
    rect(s, Inches(0.7), cy, Inches(11.9), Inches(0.95), BG_PANEL)
    txt(s, Inches(9.6), Emu(int(cy)+int(Pt(8))), Inches(2.8), Inches(0.5), t, size=18, color=ACCENT, bold=True, rtl=False, align=PP_ALIGN.RIGHT)
    txt(s, Inches(1.0), Emu(int(cy)+int(Pt(10))), Inches(8.4), Inches(0.7), d, size=15, color=WHITE)
txt(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.5),
    "التقارير السوقية الربع سنوية = جذب أعلى القمع + تعزيز للعلامة.", size=14, color=MUTED, align=PP_ALIGN.CENTER)

# ============ SLIDE 12 — Competition ============
s = new_slide(12)
header(s, "المنافسة · COMPETITIVE LANDSCAPE", "نملأ فراغاً لا يخدمه أحد بعمق")
# 2x2 matrix
mx, my, mw, mh = Inches(4.6), Inches(2.2), Inches(5.0), Inches(4.2)
rect(s, mx, my, mw, mh, BG_PANEL)
# axes labels
txt(s, mx, Emu(int(my)+int(mh)+int(Pt(4))), mw, Inches(0.4), "قوة الذكاء الاصطناعي ←", size=12, color=MUTED, align=PP_ALIGN.CENTER, rtl=False)
txt(s, Emu(int(mx)-int(Inches(2.2))), Emu(int(my)+int(mh)//2-int(Pt(10))), Inches(2.0), Inches(0.4), "العمق التخصصي ←", size=12, color=MUTED, align=PP_ALIGN.RIGHT)
# TenderPilot dot top-right
dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(mx)+int(mw)-int(Inches(1.4))), Emu(int(my)+int(Inches(0.5))), Inches(0.9), Inches(0.9))
dot.fill.solid(); dot.fill.fore_color.rgb = ACCENT; dot.line.fill.background(); dot.shadow.inherit=False
txt(s, Emu(int(mx)+int(mw)-int(Inches(1.55))), Emu(int(my)+int(Inches(1.45))), Inches(1.2), Inches(0.4), "TenderPilot", size=12, color=ACCENT, bold=True, align=PP_ALIGN.CENTER, rtl=False)
# alternatives
alts = ["مكاتب استشارية: بطيئة ومكلفة", "تنبيهات بدائية: بلا تحليل", "فرق داخلية + Word/Excel", "أدوات أجنبية عامة: لا تفهم اعتماد/العربية/الامتثال"]
bullets(s, Inches(0.7), 2.6, Inches(3.6), [("✕", a, RED) for a in alts], size=13, gap=0.7)

# ============ SLIDE 13 — Architecture ============
s = new_slide(13)
header(s, "المعمارية والأمن · ARCHITECTURE", "مبنية لثقة المؤسسات والحكومة")
layers = [
    ("Web App (Next.js)", "واجهة عربية أولاً · RTL · Dark/Light"),
    ("Agent Orchestrator", "LangGraph / Semantic Kernel + Human-in-the-Loop"),
    ("RAG + LLM Router", "Azure AI Search (Hybrid) + Citations إلزامية · توجيه ذكي للتكلفة"),
    ("Data Layer", "ADLS Gen2 (Medallion) · PostgreSQL · Vector Store · Purview"),
    ("Azure (KSA Region)", "Multi-tenant · HA/DR · توطين بيانات سيادي"),
]
for i, (t, d) in enumerate(layers):
    cy = Inches(2.2 + i*0.92)
    rect(s, Inches(0.7), cy, Inches(11.9), Inches(0.78), BG_PANEL if i % 2 else BG_DARK, line=ACCENT2)
    txt(s, Inches(9.4), Emu(int(cy)+int(Pt(8))), Inches(3.0), Inches(0.5), t, size=16, color=ACCENT, bold=True, rtl=False, align=PP_ALIGN.RIGHT)
    txt(s, Inches(0.9), Emu(int(cy)+int(Pt(10))), Inches(8.4), Inches(0.5), d, size=13, color=WHITE)

# ============ SLIDE 14 — Compliance ============
s = new_slide(14)
header(s, "الامتثال · COMPLIANCE & SECURITY", "الامتثال ميزة بيع لا عبء")
badges = ["PDPL", "ISO/IEC 27001", "NCA ECC", "SAMA"]
x = Inches(0.7); w = Inches(2.85); gap = Inches(0.18)
for i, b in enumerate(badges):
    cx = Emu(int(x) + i*(int(w)+int(gap)))
    rect(s, cx, Inches(2.4), w, Inches(1.3), BG_PANEL, line=ACCENT)
    txt(s, cx, Inches(2.85), w, Inches(0.5), b, size=18, color=ACCENT, bold=True, align=PP_ALIGN.CENTER, rtl=False)
feats = [
    "خيار No-training على بيانات العميل · عزل Multi-tenant صارم",
    "SSO / RBAC / MFA · تشفير AES-256 و TLS 1.2+ · Key Vault (HSM)",
    "Audit Log غير قابل للتعديل · SIEM (Sentinel) · اختبارات اختراق دورية",
]
bullets(s, Inches(0.7), 4.2, Inches(11.9), [("●", f, ACCENT2) for f in feats], size=16, gap=0.7)

# ============ SLIDE 15 — Roadmap ============
s = new_slide(15)
header(s, "خارطة الطريق · ROADMAP", "من المنتج إلى المنصة الإقليمية")
phases = [
    ("سنة 1 · Foundation", "Discovery · Tender Analyzer · Bid Writer · Opportunity Score · تكامل اعتماد"),
    ("سنة 2 · Intelligence", "Win Probability (ML) · Competitor · Contract · Compliance · Enterprise"),
    ("سنة 3 · Platform & Region", "API/Marketplace · Market Intelligence · توسّع GCC (UAE/Qatar)"),
]
x = Inches(0.7); w = Inches(3.83); gap = Inches(0.2)
for i, (t, d) in enumerate(phases):
    cx = Emu(int(x) + i*(int(w)+int(gap)))
    rect(s, cx, Inches(2.6), w, Inches(3.0), BG_PANEL)
    rect(s, cx, Inches(2.6), w, Inches(0.7), ACCENT)
    txt(s, cx, Inches(2.7), w, Inches(0.5), t, size=16, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Emu(int(cx)+int(Pt(8))), Inches(3.6), Emu(int(w)-int(Pt(16))), Inches(1.8), d, size=14, color=WHITE, align=PP_ALIGN.CENTER)

# ============ SLIDE 16 — Financials ============
s = new_slide(16)
header(s, "النموذج المالي · FINANCIALS", "مسار واضح إلى ARR بمئات الملايين")
# simple ARR bar chart
data = [("سنة 1", 3.6), ("سنة 2", 17.1), ("سنة 3", 49.5), ("سنة 4", 114), ("سنة 5", 220)]
base_y = Inches(5.8); max_h = Inches(3.0); maxv = 220
x0 = Inches(1.5); bw = Inches(1.4); step = Inches(2.1)
for i, (lbl, v) in enumerate(data):
    h = Emu(int(max_h * (v / maxv)))
    cx = Emu(int(x0) + i*int(step))
    by = Emu(int(base_y) - int(h))
    rect(s, cx, by, bw, h, ACCENT if i < 4 else GREEN)
    txt(s, cx, Emu(int(by)-int(Inches(0.4))), bw, Inches(0.4), f"{v}M", size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER, rtl=False)
    txt(s, cx, Emu(int(base_y)+int(Pt(4))), bw, Inches(0.4), lbl, size=13, color=MUTED, align=PP_ALIGN.CENTER)
txt(s, Inches(1.0), Inches(2.0), Inches(11.5), Inches(0.5),
    "ARR (مليون ريال) · Gross Margin 62% → 82% · EBITDA إيجابي في السنة الخامسة (+17M)",
    size=15, color=ACCENT2, align=PP_ALIGN.CENTER)

# ============ SLIDE 17 — Unit Economics ============
s = new_slide(17)
header(s, "اقتصاد الوحدة · UNIT ECONOMICS", "نمو فعّال رأسمالياً")
metrics = [
    ("LTV / CAC", "4.3x → 19x", "(سنة 1 → سنة 5)"),
    ("NRR", "115–125%", "صافي احتفاظ الإيراد"),
    ("Gross Margin", "حتى 82%", "في السنة الخامسة"),
]
x = Inches(0.9); w = Inches(3.7); gap = Inches(0.3)
for i, (t, big, d) in enumerate(metrics):
    cx = Emu(int(x) + i*(int(w)+int(gap)))
    rect(s, cx, Inches(2.6), w, Inches(3.0), BG_PANEL)
    txt(s, cx, Inches(2.9), w, Inches(0.5), t, size=16, color=MUTED, align=PP_ALIGN.CENTER, rtl=False)
    txt(s, cx, Inches(3.7), w, Inches(0.9), big, size=36, color=ACCENT, bold=True, align=PP_ALIGN.CENTER, rtl=False)
    txt(s, cx, Inches(4.8), w, Inches(0.6), d, size=14, color=WHITE, align=PP_ALIGN.CENTER)

# ============ SLIDE 18 — Team ============
s = new_slide(18)
header(s, "الفريق · TEAM", "خبرة مدمجة: منتج + AI + سوق حكومي")
roles = [
    ("المؤسسون", "رؤية المنتج وخبرة السوق الحكومي"),
    ("Product & AI", "بناء المنصة والوكلاء ونماذج التنبؤ"),
    ("GTM", "مبيعات وشراكات وتطوير أعمال"),
    ("مستشارون", "خبرة مشتريات حكومية / اعتماد"),
]
x = Inches(0.7); w = Inches(2.85); gap = Inches(0.18)
for i, (t, d) in enumerate(roles):
    cx = Emu(int(x) + i*(int(w)+int(gap)))
    rect(s, cx, Inches(2.5), w, Inches(3.0), BG_PANEL)
    ph = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(cx)+int(w)//2-int(Inches(0.55))), Inches(2.8), Inches(1.1), Inches(1.1))
    ph.fill.solid(); ph.fill.fore_color.rgb = BG_DARK; ph.line.color.rgb = ACCENT; ph.line.width=Pt(1.5); ph.shadow.inherit=False
    txt(s, cx, Inches(4.1), w, Inches(0.5), t, size=17, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Emu(int(cx)+int(Pt(8))), Inches(4.7), Emu(int(w)-int(Pt(16))), Inches(0.7), d, size=13, color=WHITE, align=PP_ALIGN.CENTER)

# ============ SLIDE 19 — The Ask ============
s = new_slide(19)
header(s, "الطلب الاستثماري · THE ASK", "تمويل لتسريع الريادة")
rect(s, Inches(0.7), Inches(2.3), Inches(5.8), Inches(3.3), BG_PANEL)
txt(s, Inches(0.9), Inches(2.5), Inches(5.4), Inches(0.5), "الجولة", size=16, color=ACCENT, bold=True, align=PP_ALIGN.RIGHT)
txt(s, Inches(0.9), Inches(3.1), Inches(5.4), Inches(0.9),
    [[("8M ريال", {'size': 34, 'color': WHITE, 'bold': True}), ("  Seed ", {'size': 18, 'color': ACCENT})]],
    align=PP_ALIGN.RIGHT, rtl=False)
txt(s, Inches(0.9), Inches(4.0), Inches(5.4), Inches(0.5), "ثم Series A: 40–55M ريال للتوسّع والتعمّق", size=15, color=MUTED, align=PP_ALIGN.RIGHT)
txt(s, Inches(0.9), Inches(4.7), Inches(5.4), Inches(0.6), "بناء المنتج + أول 100–150 عميل", size=15, color=WHITE, align=PP_ALIGN.RIGHT)
# use of funds
rect(s, Inches(6.8), Inches(2.3), Inches(5.8), Inches(3.3), BG_PANEL)
txt(s, Inches(7.0), Inches(2.5), Inches(5.4), Inches(0.5), "استخدام الأموال", size=16, color=ACCENT, bold=True, align=PP_ALIGN.RIGHT)
funds = [("منتج / AI", "50%", ACCENT), ("مبيعات وتسويق", "35%", ACCENT2), ("امتثال / عمليات", "15%", GREEN)]
for i, (t, p, col) in enumerate(funds):
    cy = Inches(3.2 + i*0.7)
    rect(s, Inches(7.0), cy, Inches(5.4), Inches(0.5), BG_DARK)
    rect(s, Emu(int(Inches(12.4))-int(Inches(5.4)*(int(p[:-1])/100))), cy, Emu(int(Inches(5.4)*(int(p[:-1])/100))), Inches(0.5), col)
    txt(s, Inches(7.1), Emu(int(cy)+int(Pt(4))), Inches(5.2), Inches(0.4),
        [[(p + "   ", {'size': 14, 'color': BG_DARK, 'bold': True}), (t, {'size': 14, 'color': BG_DARK, 'bold': True})]],
        align=PP_ALIGN.RIGHT)

# ============ SLIDE 20 — Vision & Exit ============
s = new_slide(20, decorate=True)
rect(s, Inches(0), Inches(0), Inches(0.18), SH, ACCENT)
txt(s, Inches(1.0), Inches(1.4), Inches(11.3), Inches(0.5), "الرؤية والخروج · VISION & EXIT", size=16, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1.0), Inches(2.1), Inches(11.3), Inches(0.9),
    "نظام تشغيل المناقصات للخليج", size=38, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1.0), Inches(3.3), Inches(11.3), Inches(0.6),
    "كل شركة في GCC تدير مناقصاتها عبر TenderPilot AI", size=20, color=ACCENT2, align=PP_ALIGN.CENTER)
rect(s, Inches(2.5), Inches(4.3), Inches(8.3), Inches(1.3), BG_PANEL)
txt(s, Inches(2.7), Inches(4.5), Inches(7.9), Inches(0.5),
    "Exit: استحواذ استراتيجي (SaaS/ERP عالمي · مجموعة تقنية حكومية · استشارات كبرى) أو IPO إقليمي",
    size=15, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, Inches(2.7), Inches(5.05), Inches(7.9), Inches(0.5),
    "التقييم المستهدف سنة 5: 1.1–1.3 مليار ريال (5–6x ARR)", size=16, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.6),
    "Let's build it.", size=22, color=WHITE, bold=True, align=PP_ALIGN.CENTER, rtl=False)

prs.save("/home/user/n8n-docs/TenderPilot_AI_Pitch_Deck.pptx")
print("SAVED:", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
