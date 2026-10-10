"""Proposal → Word document (right-to-left Arabic).

Placeholders ([يُستكمل: ...]) are highlighted so they can't slip through to
submission unnoticed.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .models import Proposal, Risk
from .report import LEVEL, LEVEL_RANK, STATUS
from .writer import PLACEHOLDER

FONT = "Arial"
BRAND = RGBColor(0x0B, 0x4F, 0x6C)


# OOXML is order-sensitive: Word rejects or "repairs" files whose child
# elements are out of schema order, so new elements go before their successors.
_PPR_AFTER_BIDI = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
_TBLPR_AFTER_BIDI = (
    "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing",
    "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
    "w:tblCaption", "w:tblDescription", "w:tblPrChange",
)
_RPR_AFTER_SZCS = (
    "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl",
    "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
)


def _insert(parent, tag: str, successors: tuple[str, ...], **attrs: str):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), v)
    for child in parent:
        if child.tag in {qn(t) for t in successors}:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def _rtl(paragraph) -> None:
    _insert(paragraph._p.get_or_add_pPr(), "w:bidi", _PPR_AFTER_BIDI)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _run(paragraph, text: str, *, bold: bool = False, size: int | None = None, highlight: bool = False):
    run = paragraph.add_run(text)
    font = run.font
    font.rtl = True
    font.bold = font.cs_bold = bold or None
    rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size:
        font.size = Pt(size)
        _insert(run._r.rPr, "w:szCs", _RPR_AFTER_SZCS, **{"w:val": str(size * 2)})
    if highlight:
        font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _text(container, text: str, *, bold: bool = False, size: int | None = None):
    """Paragraph with placeholders highlighted."""
    p = container.add_paragraph()
    _rtl(p)
    pos = 0
    for m in PLACEHOLDER.finditer(text):
        if m.start() > pos:
            _run(p, text[pos : m.start()], bold=bold, size=size)
        _run(p, m.group(0), bold=True, size=size, highlight=True)
        pos = m.end()
    if pos < len(text):
        _run(p, text[pos:], bold=bold, size=size)
    return p


def _heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    _rtl(p)
    run = _run(p, text, bold=True, size={0: 26, 1: 16, 2: 13}[level])
    run.font.color.rgb = BRAND


def _table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _insert(table._tbl.tblPr, "w:bidiVisual", _TBLPR_AFTER_BIDI)
    for cell, h in zip(table.rows[0].cells, headers):
        cell.paragraphs[0].text = ""
        _rtl(cell.paragraphs[0])
        _run(cell.paragraphs[0], h, bold=True, size=10)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            _text(cell, value, size=10)
    doc.add_paragraph()


def _risk_rows(risks: list[Risk]) -> list[list[str]]:
    ordered = sorted(risks, key=lambda k: -LEVEL_RANK[k.likelihood] * LEVEL_RANK[k.impact])
    return [[k.title, LEVEL[k.likelihood], LEVEL[k.impact], k.mitigation] for k in ordered]


def build_docx(proposal: Proposal, risks: list[Risk]) -> bytes:
    doc = Document()
    d = proposal.draft

    _heading(doc, "العرض الفني", 0)
    _text(doc, proposal.tender_title, bold=True, size=16)
    _text(doc, f"الجهة: {proposal.issuing_entity}")
    if proposal.tender_reference:
        _text(doc, f"رقم المنافسة: {proposal.tender_reference}")
    _text(doc, f"مقدّم من: {proposal.company_name}")
    doc.add_page_break()

    _heading(doc, "الملخص التنفيذي")
    for para in d.executive_summary:
        _text(doc, para)

    for section in d.sections:
        _heading(doc, section.title)
        for para in section.paragraphs:
            _text(doc, para)

    if d.implementation_plan:
        _heading(doc, "خطة التنفيذ")
        _table(
            doc,
            ["المرحلة", "المدة (أسابيع)", "الأنشطة", "المخرجات"],
            [
                [p.name, str(p.duration_weeks) if p.duration_weeks else "—",
                 "؛ ".join(p.activities), "؛ ".join(p.deliverables)]
                for p in d.implementation_plan
            ],
        )

    if d.staffing:
        _heading(doc, "فريق العمل")
        _table(
            doc,
            ["الدور", "العدد", "المسؤوليات", "المؤهلات"],
            [[r.role, str(r.count) if r.count else "—", r.responsibilities, r.qualifications] for r in d.staffing],
        )

    if risks:
        _heading(doc, "مصفوفة المخاطر")
        _table(doc, ["الخطر", "الاحتمال", "الأثر", "إجراء التخفيف"], _risk_rows(risks))

    if d.quality_plan:
        _heading(doc, "خطة الجودة")
        _table(
            doc,
            ["المجال", "إجراء الضبط", "التكرار", "مؤشر الأداء"],
            [[q.area, q.control, q.frequency, q.kpi] for q in d.quality_plan],
        )

    _heading(doc, "مصفوفة الامتثال")
    _table(
        doc,
        ["رقم", "المتطلب", "إلزامي", "موضع الاستجابة في العرض"],
        [
            [c.requirement_id, c.text, "نعم" if c.mandatory else "لا", "، ".join(c.sections) or "[يُستكمل: لم يُغطَّ]"]
            for c in proposal.compliance
        ],
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def review_notes(proposal: Proposal) -> str:
    """Internal checklist for the bid team — kept out of the client document."""
    lines = [f"# مراجعة داخلية: {proposal.tender_title}", ""]
    if proposal.uncovered_mandatory:
        lines += ["## متطلبات إلزامية لم يغطّها العرض", ""]
        lines += [f"- [ ] {u}" for u in proposal.uncovered_mandatory] + [""]
    gaps = [c for c in proposal.compliance if c.status in ("partial", "unmet", "unknown")]
    if gaps:
        lines += ["## متطلبات تحتاج تأكيداً (حسب تقييم الملاءمة)", ""]
        lines += [f"- [ ] {c.requirement_id} ({STATUS[c.status]}): {c.text}" for c in gaps] + [""]
    if proposal.placeholders:
        lines += ["## معلومات يجب استكمالها (مظللة بالأصفر في الملف)", ""]
        lines += [f"- [ ] {p}" for p in proposal.placeholders] + [""]
    if len(lines) == 2:
        lines.append("لا توجد ملاحظات مفتوحة.")
    return "\n".join(lines) + "\n"
