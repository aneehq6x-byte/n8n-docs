import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from tenderpilot.api import create_app
from tenderpilot.docx_export import build_docx, review_notes
from tenderpilot.models import (
    AnalysisResult,
    Phase,
    ProposalDraft,
    ProposalSection,
    QualityControl,
    StaffRole,
    VerificationReport,
)
from tenderpilot.writer import BidWriter, WriterError

from conftest import FakeLLM, fit

UNVERIFIED = VerificationReport(text_layer_available=False, checked=0, verified=0, checks=[])


@pytest.fixture
def analysis(extraction):
    return AnalysisResult(
        extraction=extraction,
        verification=UNVERIFIED,
        fit=fit(R1="met", R2="met", R3="partial", R4="unknown"),
    )


@pytest.fixture
def draft():
    # R4 (mandatory) deliberately left unaddressed.
    return ProposalDraft(
        executive_summary=["نقدّم عرضنا لتشغيل وصيانة المباني بخبرة [يُستكمل: عدد سنوات الخبرة] سنة."],
        sections=[
            ProposalSection(key="methodology", title="منهجية العمل",
                            paragraphs=["نطبّق خطة صيانة وقائية شهرية."], addresses=["R1", "R3"]),
            ProposalSection(key="quality", title="الجودة",
                            paragraphs=["نلتزم بنظام ISO 9001 [يُستكمل: رقم الشهادة وتاريخ انتهائها]."],
                            addresses=["R2"]),
        ],
        implementation_plan=[Phase(name="التجهيز", duration_weeks=4, activities=["حصر الأصول"], deliverables=["سجل الأصول"])],
        staffing=[StaffRole(role="مدير مشروع", count=1, responsibilities="الإشراف",
                            qualifications="PMP", requirement_ids=["R3"])],
        quality_plan=[QualityControl(area="الصيانة", control="تدقيق", frequency="شهري", kpi="نسبة الإنجاز")],
    )


def make_writer(draft):
    llm = FakeLLM({"write": draft})
    return BidWriter(settings=llm.settings, llm=llm), llm  # type: ignore[arg-type]


def test_compliance_matrix_and_gaps(analysis, profile, draft):
    writer, llm = make_writer(draft)
    proposal = writer.write(analysis, profile, notes="سنتعاقد مع مقاول باطن للتكييف.")

    prompt = llm.calls[0]["content"][0]["text"]
    assert "<bid_team_guidance>" in prompt and "<assessment>" in prompt

    rows = {c.requirement_id: c for c in proposal.compliance}
    assert rows["R3"].sections == ["منهجية العمل", "فريق العمل"]
    assert rows["R3"].status == "partial"
    assert proposal.uncovered_mandatory == ["R4: سجل تجاري ساري"]
    assert proposal.placeholders == ["عدد سنوات الخبرة", "رقم الشهادة وتاريخ انتهائها"]


def test_requires_requirements(analysis, profile, draft):
    analysis.extraction.requirements = []
    writer, _ = make_writer(draft)
    with pytest.raises(WriterError):
        writer.write(analysis, profile)


def test_docx_is_rtl_and_highlights_placeholders(analysis, profile, draft):
    writer, _ = make_writer(draft)
    proposal = writer.write(analysis, profile)
    doc = Document(io.BytesIO(build_docx(proposal, analysis.extraction.risks)))

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "العرض الفني" in text and "منهجية العمل" in text
    body = doc.paragraphs[-1]._p.getparent()
    assert body.xpath(".//w:bidi") and body.xpath(".//w:bidiVisual")
    highlighted = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color]
    assert "[يُستكمل: عدد سنوات الخبرة]" in highlighted
    # Compliance matrix flags the uncovered mandatory requirement.
    last_table = doc.tables[-1]
    r4 = next(row for row in last_table.rows if row.cells[0].text == "R4")
    assert "لم يُغطَّ" in r4.cells[3].text


def test_review_notes_list_open_items(analysis, profile, draft):
    writer, _ = make_writer(draft)
    notes = review_notes(writer.write(analysis, profile))
    assert "R4: سجل تجاري ساري" in notes
    assert "R3 (جزئي)" in notes and "R4 (غير معروف)" in notes
    assert "- [ ] رقم الشهادة وتاريخ انتهائها" in notes


def test_proposal_api(analysis, profile, draft):
    writer, _ = make_writer(draft)
    client = TestClient(create_app(writer=writer))
    body = {"analysis": analysis.model_dump(mode="json"), "profile": profile.model_dump(mode="json")}

    r = client.post("/v1/proposal", json=body)
    assert r.status_code == 200 and r.json()["uncovered_mandatory"]

    r = client.post("/v1/proposal?format=docx", json=body)
    assert r.status_code == 200 and r.content[:2] == b"PK"

    r = client.post("/v1/proposal?format=review", json=body)
    assert "مراجعة داخلية" in r.text

    body["analysis"]["extraction"]["requirements"] = []
    assert client.post("/v1/proposal", json=body).status_code == 422


# Word rejects or "repairs" documents whose property children are out of
# ECMA-376 schema order (LibreOffice is lenient, so it wouldn't catch this).
SCHEMA_ORDER = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint "
    "noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect "
    "bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers "
    "pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE "
    "autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents "
    "suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr "
    "sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc "
    "tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription "
    "tblPrChange",
}


def test_docx_property_elements_follow_schema_order(analysis, profile, draft):
    import zipfile

    from lxml import etree

    writer, _ = make_writer(draft)
    data = build_docx(writer.write(analysis, profile), analysis.extraction.risks)
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for name, order in SCHEMA_ORDER.items():
        order = order.split()
        for el in root.iter(ns + name):
            tags = [c.tag.removeprefix(ns) for c in el]
            assert all(t in order for t in tags), (name, tags)
            positions = [order.index(t) for t in tags]
            assert positions == sorted(positions) and len(set(tags)) == len(tags), (name, tags)
